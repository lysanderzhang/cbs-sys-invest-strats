"""Read-only checks of source provenance, saved workbook and sensitivity ledgers."""
from pathlib import Path
import hashlib
import json
import numpy as np
import pandas as pd
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/validation'
payload=json.loads((OUT/'workbook_data.json').read_text())
source=openpyxl.load_workbook(payload['source'],data_only=True,read_only=True)
cached={name:{cell.coordinate:cell.value for row in source[name] for cell in row if cell.value is not None} for name in ['AV.Clean','BBRG.Daily']}
source_checks=0
def cell_equal(address,value):
    global source_checks
    sheet,cell=address.split('!')
    assert abs(float(cached[sheet][cell])-value)<1e-9,(address,value,cached[sheet][cell])
    source_checks+=1

for r in payload['noise']:
    cell_equal(r['open_cell'],r['open_proxy'])
    if r['available']: cell_equal(r['price_cell'],r['historical_price'])
for r in payload['volatility']:
    cell_equal(r['previous_close_cell'],r['previous_close']);cell_equal(r['close_cell'],r['close'])
for r in payload['decisions']:
    for c,v in [('interval_price_cell','interval_price'),('volume_cell','volume'),('execution_price_cell','price')]:cell_equal(r[c],r[v])
for r in payload['sizing']:
    cell_equal(r['open_source_cell'],r['open_proxy']);cell_equal(r['previous_close_cell'],r['previous_close'])

book=openpyxl.load_workbook(OUT/'validation_tables.xlsx',data_only=True,read_only=True)
cell_checks=json.loads((OUT/'workbook_formula_checks.json').read_text())
saved_values={name:{cell.coordinate:cell.value for row in book[name] for cell in row if cell.value is not None} for name in book.sheetnames}
for c in cell_checks:
    actual=saved_values[c['sheet']][c['cell']]
    assert isinstance(actual,(int,float)) and abs(actual-c['expected'])<=c['tolerance'],(c,actual)
source.close();book.close()

ledger_checks=[]
for case in payload['case_definitions']:
    for model in 'ABC':
        prefix=OUT/'scenarios'/case['name']
        day=pd.read_csv(prefix/f'{model}_daily.csv')
        tr=pd.read_csv(prefix/f'{model}_trades.csv')
        ev=pd.read_csv(prefix/f'{model}_events.csv')
        dec=pd.read_csv(prefix/f'{model}_decisions.csv')
        assert len(day)==168
        np.testing.assert_allclose(tr.net_pnl.sum(),day.net_pnl.sum(),rtol=0,atol=1e-7)
        np.testing.assert_allclose(ev.net_pnl.sum(),day.net_pnl.sum(),rtol=0,atol=1e-7)
        np.testing.assert_allclose(dec.interval_net_pnl.sum(),day.net_pnl.sum(),rtol=0,atol=1e-7)
        np.testing.assert_allclose(day.gross_pnl-day.costs,day.net_pnl,rtol=0,atol=1e-7)
        np.testing.assert_allclose(ev.commission,ev.shares*case['commission'],rtol=0,atol=1e-9)
        np.testing.assert_allclose(ev.slippage,ev.shares*case['slippage'],rtol=0,atol=1e-9)
        assert (dec.groupby('date').held_after.last()==0).all()
        for special in ['2025-07-03','2025-11-28']:
            z=ev[ev.date==special]
            assert (pd.to_datetime(z.timestamp).dt.time<=pd.Timestamp('13:00').time()).all()
        if case['delayed']:
            z=ev[~ev.is_close]
            assert ((pd.to_datetime(z.timestamp)-pd.to_datetime(z.signal_timestamp))==pd.Timedelta(minutes=30)).all()
        if case['full_history_only']:
            z=dec[(dec.noise_count<14)&~dec.is_close]
            assert (z.held_before==z.held_after).all()
        ledger_checks.append(dict(case=case['name'],model=model,passed=True,trades=len(tr)))

before=json.loads((OUT/'baseline_hashes_before.json').read_text())
for name,digest in before.items():
    assert hashlib.sha256(Path(name).read_bytes()).hexdigest()==digest,name
result=dict(source_cell_comparisons=source_checks,saved_workbook_formula_caches_checked=len(cell_checks),
    scenario_ledgers_checked=len(ledger_checks),scenario_checks=ledger_checks,
    protected_files_unchanged=len(before),native_excel_tested=False,all_passed=True)
(OUT/'final_verification.json').write_text(json.dumps(result,indent=2))
print(json.dumps({k:v for k,v in result.items() if k!='scenario_checks'},indent=2))
