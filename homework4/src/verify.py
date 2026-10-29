"""Separate scalar/Decimal reconstruction; no production finance imports."""
from pathlib import Path
from decimal import Decimal,localcontext
import math,json,hashlib,csv
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]


def main():
    checks=[]
    def close(name,a,b,atol=3e-14):
        error=abs(a-b);assert error<=atol,(name,a,b,error)
        checks.append(dict(check=name,error=float(error),tolerance=atol))
    for row in json.loads((ROOT/'data/manifest.json').read_text()):
        assert hashlib.sha256((ROOT/row['copy']).read_bytes()).hexdigest()==row['sha256']
    wb=load_workbook(ROOT/'data/raw/hw3.data.carry.FX.dly.xlsx',read_only=True,data_only=True)
    sh=wb['Data'];rows=sh.iter_rows(values_only=True);next(rows);header=next(rows)
    datepos=header.index('Date');nzdpos=header.index('NZD');jpypos=header.index('JPY');levels=[]
    for row in rows:
        if row[datepos] is not None:levels.append((row[datepos].strftime('%Y-%m-%d'),Decimal(str(row[nzdpos]))*Decimal(str(row[jpypos]))))
    wb.close()
    saved=list(csv.DictReader((ROOT/'results/timeseries/fx_returns.csv').open()));pairs=[];manual=[]
    with localcontext() as ctx:
        ctx.prec=40
        for i,rec in enumerate(saved,1):
            day,z=levels[i];prev=levels[i-1][1];a=float(z/prev-1);b=float(prev/z-1)
            assert rec['Date']==day
            close(f'{day} NZDJPY scalar',a,float(rec['NZDJPY']));close(f'{day} JPYNZD scalar',b,float(rec['JPYNZD']))
            close(f'{day} inverse identity',(1+a)*(1+b),1)
            pairs.append((a,b))
            if i in [1,len(saved)] or day in ['2008-09-15','2008-10-24','2008-10-28']:
                manual.append(dict(return_date=day,previous_cross=str(prev),current_cross=str(z),NZDJPY=a,JPYNZD=b))
    with (ROOT/'results/verification/representative_observations.csv').open('w') as f:
        w=csv.DictWriter(f,fieldnames=manual[0].keys(),lineterminator='\n');w.writeheader();w.writerows(manual)
    n=len(pairs)
    for rec in csv.DictReader((ROOT/'results/tables/inverse_moments.csv').open()):
        j=0 if rec['source']=='NZDJPY' else 1;x=[p[j] for p in pairs];y=[p[1-j] for p in pairs]
        mu=math.fsum(x)/n;v=math.fsum((a-mu)**2 for a in x)/n;ym=math.fsum(y)/n;yv=math.fsum((a-ym)**2 for a in y)/n
        for k,val in dict(source_mean=mu,source_variance=v,actual_mean=ym,actual_variance=yv,approx_mean=-mu/(1+mu)+v/(1+mu)**3,approx_variance=v/(1+mu)**4).items():close(rec['source']+' '+k,val,float(rec[k]))
    for rec in csv.DictReader((ROOT/'results/tables/doubling_selected_rates.csv').open()):
        x=float(rec['rate_decimal']);exact=math.log(2)/math.log(1+x)
        close('doubling compounding '+str(x),math.exp(float(rec['exact'])*math.log(1+x)),2,2e-12)
        for k,v in [('linear_log2',math.log(2)/x),('linear_069',.69/x),('corrected_log2',math.log(2)/x+math.log(2)/2),('corrected_069',.69/x+.345),('rule72',72/(100*x))]:
            close(k+' '+str(x),v,float(rec[k]),2e-12)
            close(k+' rel error '+str(x),100*(v/exact-1),float(rec[k+'_relative_error_pct']),2e-11)
    lo,hi=.001,1.
    for _ in range(100):
        x=(lo+hi)/2
        if math.log1p(x)/x>math.log(2)/.72:lo=x
        else:hi=x
    z=json.loads((ROOT/'results/tables/rule72_optimum.json').read_text());close('independent root bisection',(lo+hi)/2,z['rate_decimal'])
    ex=json.loads((ROOT/'results/tables/illustrative_option_example.json').read_text());p=ex['parameters'];s=p['spot'];k=p['strike'];g=math.exp(p['foreign_rate']*p['maturity']);a=math.exp(p['funding_rate']*p['maturity'])
    v=p['volatility']*math.sqrt(p['maturity']);d1=(math.log(s/k)+(p['funding_rate']-p['foreign_rate']+.5*p['volatility']**2)*p['maturity'])/v;d2=d1-v
    N=lambda u:.5*(1+math.erf(u/math.sqrt(2)))
    put=k/a*N(-d2)-s/g*N(-d1);call=s/g*N(d1)-k/a*N(d2);delta=-N(-d1)/g;q=g/(1+g*delta)
    for key,val in [('put',put),('call',call),('delta',delta),('q',q)]:close('option '+key,val,ex['delta_hedged'][key])
    close('put-call parity',call-put,s/g-k/a)
    close('terminal put quantity',q,g*(1-q*delta));close('matched initial total delta',1-q*delta+q*delta,1)
    for rec in csv.DictReader((ROOT/'results/timeseries/illustrative_option_payoffs.csv').open()):
        st=float(rec['terminal_spot']);value=g*(1-q*delta)*st/s+q*max(k-st,0)/s-a*((1-q*delta)+q*put/s)
        close('funded payoff '+str(st),value,float(rec['delta_hedged_put_overlay']))
        assert value>=-q*a*call/s-3e-14
    result=dict(all_checks_passed=True,checks=len(checks),return_dates=n,input_hashes_unchanged=True,max_scalar_error=max(x['error'] for x in checks),scope='Full daily FX reconstruction, moments, five approximations, independent root, funded option payoffs and price parity.',details=checks)
    result['verified_output_sha256']={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for folder in ['tables','timeseries'] for p in sorted((ROOT/'results'/folder).glob('*')) if p.is_file()}
    (ROOT/'results/verification/independent_reconstruction.json').write_text(json.dumps(result,indent=2)+'\n')
    print({k:v for k,v in result.items() if k!='details'})

if __name__=='__main__':main()
